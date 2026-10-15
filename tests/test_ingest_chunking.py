from __future__ import annotations

import pymupdf
import pytest
from PIL import Image

from course_assistant.chunking import chunk_pages
from course_assistant.ingest import UnsupportedFormat, find_soffice, ingest_file


def make_pdf(path, pages):
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
    doc.save(path)


def test_pdf_pages_keep_text_and_images(tmp_path):
    pdf = tmp_path / "lecture.pdf"
    make_pdf(pdf, ["Gradient descent minimises the loss function step by step.",
                   "Learning rate controls the step size of each update in training."])
    info, pages = ingest_file(pdf, tmp_path / "data")
    assert info.num_pages == 2 and info.has_images
    assert [p.number for p in pages] == [1, 2]
    assert all(p.kind == "page" and p.doc_name == "lecture.pdf" for p in pages)
    assert "Gradient descent" in pages[0].text
    for p in pages:
        with Image.open(p.image_path) as im:
            assert im.width == 1280


def test_scanned_pdf_page_is_ocrd(tmp_path, fake):
    pdf = tmp_path / "scan.pdf"
    make_pdf(pdf, [""])  # no text layer, like a scan
    _, pages = ingest_file(pdf, tmp_path / "data", services=fake)
    assert pages[0].ocr and "photosynthesis" in pages[0].text


def test_markdown_splits_by_heading(tmp_path):
    md = tmp_path / "notes.md"
    md.write_text("# Intro\nHello there world.\n\n## Details\nMore text here.\n")
    _, pages = ingest_file(md, tmp_path / "data")
    assert [p.title for p in pages] == ["Intro", "Details"]
    assert all(p.kind == "section" and p.image_path is None for p in pages)


def test_html_drops_scripts_and_splits(tmp_path):
    h = tmp_path / "page.html"
    h.write_text("<html><script>var secret=1</script><h1>Syllabus</h1><p>Grading: 40% exams.</p>"
                 "<h2>Schedule</h2><p>Week 1 intro.</p></html>")
    _, pages = ingest_file(h, tmp_path / "data")
    assert [p.title for p in pages] == ["Syllabus", "Schedule"]
    assert "secret" not in " ".join(p.text for p in pages)


def test_txt_and_image(tmp_path, fake):
    t = tmp_path / "a.txt"
    t.write_text("First paragraph.\n\nSecond paragraph.")
    _, pages = ingest_file(t, tmp_path / "data")
    assert "Second paragraph" in pages[0].text

    img = tmp_path / "diagram.png"
    Image.new("RGB", (50, 50), "white").save(img)
    info, pages = ingest_file(img, tmp_path / "data", services=fake)
    assert info.kind == "image" and pages[0].image_path and pages[0].ocr


def test_pptx_text_notes_and_slide_count(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"Slide title {i + 1}"
        s.placeholders[1].text = f"Body text for slide number {i + 1}"
        s.notes_slide.notes_text_frame.text = f"Speaker note {i + 1}"
    prs.slides[0].shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1)).text_frame.text = "extra box"
    path = tmp_path / "deck.pptx"
    prs.save(path)
    info, pages = ingest_file(path, tmp_path / "data")
    assert [p.number for p in pages] == [1, 2, 3]
    assert pages[1].title == "Slide title 2" and pages[1].kind == "slide"
    assert "Speaker note 2" in pages[1].text and "extra box" in pages[0].text
    if find_soffice():
        assert all(p.image_path for p in pages)
    else:
        assert info.warnings  # text-only fallback is reported


def test_docx_fallback_or_pages(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Course policies", 1)
    d.add_paragraph("Late work loses 10% per day.")
    path = tmp_path / "syllabus.docx"
    d.save(path)
    _, pages = ingest_file(path, tmp_path / "data")
    assert any("Late work loses 10% per day" in p.text for p in pages)


def test_unsupported_format_lists_accepted(tmp_path):
    f = tmp_path / "grades.xlsx"
    f.write_bytes(b"x")
    with pytest.raises(UnsupportedFormat, match=r"\.pdf"):
        ingest_file(f, tmp_path / "data")


def test_chunks_stay_within_one_page_and_keep_offsets(tmp_path):
    md = tmp_path / "long.md"
    body_a = " ".join(f"Alpha sentence number {i}." for i in range(200))
    body_b = " ".join(f"Beta sentence number {i}." for i in range(200))
    md.write_text(f"# A\n{body_a}\n\n# B\n{body_b}\n")
    _, pages = ingest_file(md, tmp_path / "data")
    chunks = chunk_pages(pages, size=300, overlap=50)
    assert len(chunks) > 4
    by_num = {p.number: p for p in pages}
    for c in chunks:
        page = by_num[c.number]
        assert page.text[c.start:c.start + len(c.text)] == c.text  # offset is exact
        assert ("Alpha" in c.text) != ("Beta" in c.text)  # never mixes pages
        assert c.doc_name == "long.md" and c.label.startswith("Section")


def test_ocr_text_missing_from_text_layer_becomes_visual_chunks(tmp_path, fake):
    pdf = tmp_path / "slides.pdf"
    make_pdf(pdf, ["Other Prompting Techniques. Source: Walter 2024, page fifteen."])
    fake.ocr_text = ("Table 4 Chain-of-Thought prompting: the AI should slowly elaborate on how a given "
                     "response is generated. Self-consistency: generate several answers and discern "
                     "the best one. Tree-of-thought prompting explores branches.")
    _, pages = ingest_file(pdf, tmp_path / "data", services=fake)
    assert not pages[0].ocr and "Chain-of-Thought" in pages[0].visual_text
    chunks = chunk_pages(pages)
    vis = [c for c in chunks if c.field == "visual_text"]
    assert vis and vis[0].chunk_id.endswith(":v0")
    assert pages[0].visual_text[vis[0].start:vis[0].start + len(vis[0].text)] == vis[0].text


def test_ocr_that_adds_nothing_is_not_stored(tmp_path, fake):
    pdf = tmp_path / "plain.pdf"
    text = "Learning rate controls the step size of each update during model training."
    make_pdf(pdf, [text])
    fake.ocr_text = text
    _, pages = ingest_file(pdf, tmp_path / "data", services=fake)
    assert pages[0].visual_text == ""


def test_clean_ocr_flattens_tables_and_markdown():
    from course_assistant.ingest import clean_ocr

    raw = "**Table 4**\n<table><tr><td>Method</td><td>Acronym</td></tr><tr><td>Chain-of-Thought</td><td>CoT</td></tr></table>"
    assert clean_ocr(raw) == "Table 4\nMethod | Acronym\nChain-of-Thought | CoT"


def _deck_with_picture_and_video(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    gif = tmp_path / "anim.gif"
    def striped(color):  # pictures need visible content (a flat colour counts as "no content")
        im = Image.new("RGB", (100, 60), color)
        for x in range(0, 100, 10):
            im.paste((255, 255, 255), (x, 0, x + 3, 60))
        return im
    frames = [striped(c) for c in ("blue", "red", "yellow")]  # middle frame is red
    frames[0].save(gif, save_all=True, append_images=frames[1:], duration=100, loop=0)  # animated GIF
    poster = tmp_path / "poster.png"
    striped("green").save(poster)
    video = tmp_path / "clip.mp4"
    video.write_bytes(b"\\x00" * 64)

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    s1.shapes.title.text = "Slide with an animated GIF"
    s1.shapes.add_picture(str(gif), Inches(1), Inches(2), Inches(4), Inches(2.4))
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Slide with a video"
    s2.shapes.add_movie(str(video), Inches(5), Inches(2), Inches(4), Inches(2.4),
                        poster_frame_image=str(poster), mime_type="video/mp4")
    path = tmp_path / "media.pptx"
    prs.save(path)
    return path, prs


def test_blank_media_areas_are_filled_and_reported(tmp_path):
    from course_assistant.ingest import repair_pptx_renders

    deck, prs = _deck_with_picture_and_video(tmp_path)
    w = 1280
    h = int(w * prs.slide_height / prs.slide_width)
    blank = []
    for i in range(2):  # simulate LibreOffice leaving the media areas blank (all white)
        p = tmp_path / f"render{i}.png"
        Image.new("RGB", (w, h), "white").save(p)
        blank.append(str(p))
    notes, warnings = repair_pptx_renders(deck, blank)

    scale = w / prs.slide_width
    from pptx.util import Inches
    def colours(path, left_in):  # distinct colours inside the pasted area
        im = Image.open(path).convert("RGB")
        box = (int(Inches(left_in) * scale), int(Inches(2) * scale),
               int(Inches(left_in + 4) * scale), int(Inches(4.4) * scale))
        return {c for _, c in im.crop(box).getcolors(1_000_000)}
    assert any(r > 200 and g < 60 for r, g, b in colours(blank[0], 1))   # middle GIF frame (red) pasted in
    assert any(g > 100 and r < 60 for r, g, b in colours(blank[1], 5))   # video poster (green) pasted in
    assert len(warnings) == 2 and "animated GIF or video" in warnings[0]
    assert 2 in notes and "embedded video" in notes[2] and 1 not in notes


def test_flat_pictures_are_not_mistaken_for_failed_renders(tmp_path):
    """A mostly-white chart or logo legitimately renders as a near-uniform area."""
    from pptx import Presentation
    from pptx.util import Inches

    from course_assistant.ingest import repair_pptx_renders

    white = tmp_path / "white.png"
    Image.new("RGB", (100, 60), "white").save(white)
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(str(white), Inches(1), Inches(1), Inches(4))
    deck = tmp_path / "flat.pptx"
    prs.save(deck)
    render = tmp_path / "r.png"
    Image.new("RGB", (1280, 960), "white").save(render)
    assert repair_pptx_renders(deck, [str(render)])[1] == []


def test_areas_that_rendered_fine_are_left_alone(tmp_path):
    from course_assistant.ingest import repair_pptx_renders

    deck, prs = _deck_with_picture_and_video(tmp_path)
    w = 1280
    h = int(w * prs.slide_height / prs.slide_width)
    renders = []
    for i in range(2):  # a "busy" render everywhere: noise, so no area is blank
        p = tmp_path / f"busy{i}.png"
        Image.effect_noise((w, h), 80).convert("RGB").save(p)
        renders.append(str(p))
    before = [Image.open(r).tobytes() for r in renders]
    _, warnings = repair_pptx_renders(deck, renders)
    assert warnings == [] and [Image.open(r).tobytes() for r in renders] == before
